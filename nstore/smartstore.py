"""마스터 시트 → 스마트스토어 상품 일괄등록 엑셀.

스마트스토어센터에서 받은 양식 파일을 그대로 열어 헤더를 찾고 그 아래에 채웁니다.
(양식의 서식·드롭다운을 유지하기 위해 pandas 대신 openpyxl로 씁니다.)
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from .config import load_yaml
from .excel_io import normalize_header
from .titles import LEVEL_ERROR, check_title

# 최종* 칸이 비어 있으면 추천값을 씀
FALLBACKS = {"최종상품명": "추천상품명", "최종태그": "추천태그"}
NUMERIC_COLUMNS = {"판매가", "재고"}


def _value(master_col: str, raw: str):
    if master_col in NUMERIC_COLUMNS and raw:
        try:
            return int(float(str(raw).replace(",", "")))
        except ValueError:
            return raw
    return raw


def _locate_header(wb: openpyxl.Workbook, expected: list[str], scan: int = 15):
    """(시트, 헤더 행 번호, {정규화된 헤더: 열 번호}) — 기대 헤더가 가장 많이 맞는 곳."""
    expected_norm = {normalize_header(e) for e in expected}
    best = (None, None, {}, 0)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=scan):
            positions = {normalize_header(c.value): c.column for c in row if c.value is not None}
            hits = len(expected_norm & positions.keys())
            if hits > best[3]:
                best = (ws, row[0].row, positions, hits)
    return best[:3]


def export(master: pd.DataFrame, output_path: Path, template_path: Path | None, settings: dict, rules: dict) -> list[str]:
    cfg = load_yaml("smartstore_columns.yaml")
    columns: dict[str, str] = cfg.get("columns", {})
    fixed: dict[str, str] = cfg.get("fixed", {}) or {}

    df = master.copy()
    for final, recommended in FALLBACKS.items():
        df[final] = df[final].where(df[final] != "", df[recommended])

    # 내보내기 직전에 최종 상품명을 다시 검사 — 오류 있는 상품은 빼고 알려줌
    blocked = []
    keep = []
    for _, row in df.iterrows():
        title = row["최종상품명"]
        errors = [str(i) for i in check_title(title, rules, settings) if i.level == LEVEL_ERROR] if title else ["상품명 없음"]
        if errors:
            blocked.append(f"{row['자사몰상품코드'] or row['자사몰상품명']}: {', '.join(errors)}")
        else:
            keep.append(row)
    rows = pd.DataFrame(keep, columns=df.columns)

    notes = [f"스마트스토어 양식으로 {len(rows)}개 출력"]
    if blocked:
        notes.append(f"오류로 뺀 상품 {len(blocked)}개:")
        notes += ["  " + b for b in blocked[:20]]
        if len(blocked) > 20:
            notes.append(f"  … 외 {len(blocked) - 20}개")
    no_category = int((rows["카테고리ID"] == "").sum()) if len(rows) else 0
    if no_category:
        notes.append(f"카테고리ID가 빈 상품 {no_category}개 — config/items.yaml 의 category_id 를 채우세요.")

    headers = list(columns) + list(fixed)
    if template_path:
        wb = openpyxl.load_workbook(template_path)
        ws, header_row, positions = _locate_header(wb, headers)
        if ws is None:
            raise ValueError("양식에서 헤더 행을 찾지 못했습니다. config/smartstore_columns.yaml 의 헤더 이름을 확인하세요.")
        missing = [h for h in headers if normalize_header(h) not in positions]
        if missing:
            notes.append("양식에서 못 찾은 컬럼: " + ", ".join(missing))
        start = header_row + 1 + int(cfg.get("skip_rows_after_header", 0) or 0)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "일괄등록"
        positions = {}
        for i, h in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=h)
            positions[normalize_header(h)] = i
        start = 2
        notes.append("양식 파일 없이 기본 표로 출력했습니다. 실제 업로드에는 스마트스토어 양식을 지정하세요.")

    for offset, (_, product) in enumerate(rows.iterrows()):
        r = start + offset
        for header, master_col in columns.items():
            col = positions.get(normalize_header(header))
            if col:
                ws.cell(row=r, column=col, value=_value(master_col, product.get(master_col, "")))
        for header, value in fixed.items():
            col = positions.get(normalize_header(header))
            if col:
                ws.cell(row=r, column=col, value=value)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return notes
