import openpyxl
import pandas as pd

from nstore import smartstore
from nstore.master import MASTER_COLUMNS


def master_rows(*rows):
    df = pd.DataFrame([{c: "" for c in MASTER_COLUMNS} | r for r in rows])
    return df[MASTER_COLUMNS]


def make_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["상품 일괄등록 양식 — 안내문"])
    ws.append(["판매자상품코드", "카테고리코드*", "상품명 (필수)", "판매가", "상품상태", "태그", "기타"])
    ws.append(["예시", "50000000", "예시 상품", 1000, "신상품", "", ""])
    wb.save(path)


def test_export_fills_template_below_example_row(tmp_path, monkeypatch, settings, rules):
    monkeypatch.setattr(smartstore, "load_yaml", lambda _: {
        "columns": {"판매자상품코드": "자사몰상품코드", "카테고리코드": "카테고리ID", "상품명": "최종상품명",
                    "판매가": "판매가", "태그": "최종태그", "재고수량": "재고"},
        "fixed": {"상품상태": "신상품"},
        "skip_rows_after_header": 1,
    })
    template, out = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    make_template(template)
    master = master_rows(
        {"자사몰상품코드": "001", "카테고리ID": "50001234", "추천상품명": "케어메이트 성인용보행기",
         "판매가": "150,000", "추천태그": "실버카,노인보행기"},
        {"자사몰상품코드": "002", "추천상품명": "보행기", "최종상품명": "보행기 최저가"},
    )
    notes = smartstore.export(master, out, template, settings, rules)

    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[3]][:3] == ["예시", "50000000", "예시 상품"]  # 예시 행은 그대로
    assert [c.value for c in ws[4]][:6] == ["001", "50001234", "케어메이트 성인용보행기", 150000, "신상품", "실버카,노인보행기"]
    assert ws.max_row == 4  # 금지 문구 상품은 빠짐
    assert any("오류로 뺀 상품 1개" in n for n in notes)
    assert any("002" in n and "최저가" in n for n in notes)
    assert any("못 찾은 컬럼: 재고수량" in n for n in notes)


def test_export_without_template_writes_plain_sheet(tmp_path, settings, rules):
    out = tmp_path / "out.xlsx"
    master = master_rows({"자사몰상품코드": "001", "추천상품명": "케어메이트 성인용보행기"})
    notes = smartstore.export(master, out, None, settings, rules)
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=2, column=3).value == "케어메이트 성인용보행기"
    assert any("카테고리ID가 빈 상품 1개" in n for n in notes)
