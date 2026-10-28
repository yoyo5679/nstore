"""아임웹 엑셀 → 키워드 → 상품명 → 스마트스토어 엑셀 전체 흐름 (가짜 네이버 API)."""

import openpyxl
import pandas as pd
import pytest
from fakes import walker_scenario

from nstore import config, pipeline


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MASTER_PATH", tmp_path / "master.xlsx")
    monkeypatch.setattr(config, "KEYWORDS_PATH", tmp_path / "keywords.xlsx")
    monkeypatch.setattr(config, "CACHE_PATH", tmp_path / "cache.sqlite")
    monkeypatch.setattr(config, "SMARTSTORE_OUTPUT_PATH", tmp_path / "upload.xlsx")
    searchad, shopping = walker_scenario()
    monkeypatch.setattr(pipeline.SearchAdClient, "from_env", classmethod(lambda cls, cache=None: searchad))
    monkeypatch.setattr(pipeline.ShoppingClient, "from_env", classmethod(lambda cls, cache=None: shopping))
    return tmp_path


def test_full_flow(workspace):
    imweb = workspace / "imweb.xlsx"
    pd.DataFrame({
        "상품코드": ["A1"],
        "상품명": ["케어메이트 접이식 성인용 보행기 CM-200"],
        "판매가": ["150000"],
        "재고": ["10"],
        "브랜드": ["케어메이트"],
        "모델명": ["CM-200"],
    }).to_excel(imweb, index=False)

    pipeline.run_import(imweb)
    with pytest.raises(FileExistsError):
        pipeline.run_import(imweb)

    notes = pipeline.run_keywords()
    assert notes[0].startswith("성인용보행기: 관련 키워드")

    pipeline.run_titles()
    master = pd.read_excel(config.MASTER_PATH, dtype=str).fillna("")
    assert master.loc[0, "추천상품명"] == "케어메이트 CM-200 접이식 성인용보행기"
    assert master.loc[0, "검사결과"] == "통과"

    pipeline.run_export()
    ws = openpyxl.load_workbook(config.SMARTSTORE_OUTPUT_PATH).active
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["상품명"] == "케어메이트 CM-200 접이식 성인용보행기"
    assert row["판매가"] == 150000
    assert row["태그"] == "노인보행기,실버카,보행보조차"
